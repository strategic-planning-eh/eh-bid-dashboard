"""merge_water_twins.py — Youssef's 20 Oct 2026 decision: merge the two AR↔EN twins flagged by patch_water_sector_stakeholders.py.

  Red Sea Utilities Consortium            ← شركة مرافق البحر الاحمر المشتركة للطاقة (MWAN licence record, Umluj)
  Veolia Middle East Environmental Services ← فيوليا ميديل ايست للخدمات البيئية (MWAN licence record, Jubail)

Same mechanics as apply_region_review.py: the English record is kept, the Arabic name becomes "Alias (merged duplicate): …",
the higher tier / bid count carries, and — because both duplicates are MWAN-register rows — the licence block (mw), licence
activities, sites and the MWAN flags carry to the kept record so Licence Watch loses nothing. The kept record takes the
duplicate's map pin when its own was only a region centroid. Veolia's kept record moves to Direct Competitors: the licence
record (hazardous-waste treatment, 384 t/day thermal oxidation in Jubail) is EH's own service line.
The decision is written to pipeline/twin_decisions.xlsx (so the hourly sync honours it) and to pipeline/aliases.csv.

Run once: python3 pipeline/merge_water_twins.py [--map hub/EH_Stakeholder_Map_CURRENT.html] [--workbook IN.xlsx --out OUT.xlsx]
"""
import argparse, csv, json, os, re, datetime as dt
HERE = os.path.dirname(os.path.abspath(__file__))
PAIRS = [  # kept, loser, new category for the kept record (None = unchanged), reviewer note
 ('Red Sea Utilities Consortium', 'شركة مرافق البحر الاحمر المشتركة للطاقة', None,
  'Same company: Red Sea Utilities Co. (RSG subsidiary, Umluj) — SWA desalination licence and MWAN sorting/trading licence. Decided 20 Oct 2026.'),
 ('Veolia Middle East Environmental Services', 'فيوليا ميديل ايست للخدمات البيئية', 'Direct Competitors',
  'Same company: Veolia Middle East, Jubail — MWAN hazardous-waste treatment licence (thermal oxidation, 384 t/day). Direct competitor. Decided 20 Oct 2026.'),
]
TR = {'Tier 1': 1, 'GOV': 1, 'Tier 2': 2, 'Tier 3': 3, 'Tier 4': 4}
ap = argparse.ArgumentParser(); ap.add_argument('--map', default=os.path.join(HERE, '..', 'hub', 'EH_Stakeholder_Map_CURRENT.html'))
ap.add_argument('--workbook'); ap.add_argument('--out'); a = ap.parse_args()
today = dt.datetime.now(dt.timezone(dt.timedelta(hours=3))).strftime('%Y-%m-%d')

h = open(a.map, encoding='utf-8').read()
m = re.search(r'(<script id="DATA" type="application/json">)(.*?)(</script>)', h, re.S)
D = json.loads(m.group(2)); nodes = D['nodes']; by = {n['n']: n for n in nodes}
done = []
for kept, loser, cat, why in PAIRS:
    kn, ln = by.get(kept), by.get(loser)
    if kn is None: raise SystemExit(f'kept record missing: {kept}')
    notes = kn.get('notes') or ''
    notes = re.sub(r' Possible twin of the MWAN-licensed record «[^»]*» \([^)]*\) — left for the twin review\.', '', notes)
    if f'Alias (merged duplicate): {loser}' not in notes: notes = (notes + ' ' if notes else '') + f'Alias (merged duplicate): {loser}.'
    kn['notes'] = notes
    if ln is not None:
        for k in ('mw', 'sites', 'mn', 'mp'):
            if ln.get(k) and not kn.get(k): kn[k] = ln[k]
        if ln.get('act'):
            acts = [x.strip() for x in (kn.get('act') or '').split(',') if x.strip()]
            kn['act'] = ', '.join(acts + [x.strip() for x in ln['act'].split(',') if x.strip() and x.strip() not in acts])
        if ln.get('city') and not kn.get('city'): kn['city'] = ln['city']
        if TR.get(ln.get('tier'), 9) < TR.get(kn.get('tier'), 9): kn['tier_old'] = kn.get('tier'); kn['tier'] = ln['tier']
        if (ln.get('bc') or 0) > (kn.get('bc') or 0): kn['bid'] = True; kn['bc'] = ln['bc']; kn['bf'] = ln.get('bf')
        if ln.get('cf') == 'High' and kn.get('cf') != 'High': kn['cf'] = 'High'; kn['cf_basis'] = 'MWAN licence register (merged record)'
        nodes.remove(ln)
    if cat and kn.get('cat') != cat: kn['cat_old'] = kn.get('cat'); kn['cat'] = cat
    done.append((kept, loser, ln is not None))
payload = json.dumps(D, ensure_ascii=False).replace('</', '<\\/')
h = h[:m.start(2)] + payload + h[m.end(2):]
# map pins: the duplicate's precise pin replaces a region-centroid pin on the kept record, then the duplicate's entry goes
for kept, loser, _, _ in PAIRS:
    mk = re.search(r'"' + re.escape(kept) + r'":\[([^\]]*)\]', h); ml = re.search(r'"' + re.escape(loser) + r'":\[([^\]]*)\],?', h)
    if mk and ml:
        kv, lv = mk.group(1).split(','), ml.group(1).split(',')
        if len(kv[0].split('.')[-1]) <= 2 and len(lv[0].split('.')[-1]) > 2:   # centroid (2 dp) vs geocoded pin
            h = h.replace(mk.group(0), '"' + kept + '":[' + ','.join(lv[:2] + kv[2:]) + ']')
    if ml: h = h.replace(ml.group(0), '')
open(a.map, 'w', encoding='utf-8').write(h)
print('map:', [(k, 'merged' if ok else 'alias only') for k, l, ok in done], '| nodes now', len(nodes))

# ---- reviewer decision for the sync, and the alias bridge
import openpyxl
wb = openpyxl.load_workbook(os.path.join(HERE, 'twin_decisions.xlsx')); ws = wb['Possible twins']
hdr = [str(c.value or '') for c in ws[1]]; have = {(str(r[1] or '').strip(), str(r[7] or '').strip()) for r in ws.iter_rows(min_row=2, values_only=True)}
for kept, loser, _, why in PAIRS:
    if (kept, loser) in have or (loser, kept) in have: continue
    r = [None] * len(hdr); r[0] = ws.max_row; r[1] = kept; r[7] = loser; r[13] = 'High'
    r[14] = 'same company in English and Arabic (SWA water-sector chart, Oct 2026)'; r[15] = 'Merge → keep A'; r[16] = why
    ws.append(r)
wb.save(os.path.join(HERE, 'twin_decisions.xlsx'))
p = os.path.join(HERE, 'aliases.csv'); rows = list(csv.DictReader(open(p, encoding='utf-8-sig')))
known = {(r['tracker_name'], r['map_name']) for r in rows}
for kept, loser, _, _ in PAIRS:
    if (loser, kept) not in known: rows.append({'tracker_name': loser, 'map_name': kept, 'method': 'reviewer merge', 'status': 'approved', 'added': today})
with open(p, 'w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=['tracker_name', 'map_name', 'method', 'status', 'added']); w.writeheader(); w.writerows(rows)
print('twin_decisions.xlsx and aliases.csv updated')

# ---- workbook
if a.workbook:
    wb = openpyxl.load_workbook(a.workbook); ws = wb['Stakeholders']
    wc = {str(c.value or ''): i + 1 for i, c in enumerate(ws[1])}
    rows = {str(ws.cell(i, wc['Name']).value or '').strip(): i for i in range(2, ws.max_row + 1)}
    for kept, loser, cat, why in PAIRS:
        kr, lr = rows.get(kept), rows.get(loser)
        if kr is None: print('workbook: kept row missing', kept); continue
        c = ws.cell(kr, wc['Notes']); old = str(c.value or '')
        old = re.sub(r' Possible twin of the MWAN-licensed record «[^»]*» \([^)]*\) — left for the twin review\.', '', old)
        if f'Alias (merged duplicate): {loser}' not in old: old = (old + ' ' if old else '') + f'Alias (merged duplicate): {loser}.'
        c.value = old
        if cat: ws.cell(kr, wc['Category']).value = cat
        if lr is not None:
            for col in ('Licence Activities', 'MWAN Expiry', 'Days to Expiry', 'Licensed Cities'):
                if col in wc and ws.cell(lr, wc[col]).value and not ws.cell(kr, wc[col]).value: ws.cell(kr, wc[col]).value = ws.cell(lr, wc[col]).value
    for kept, loser, _, _ in PAIRS:
        lr = rows.get(loser)
        if lr: ws.delete_rows(lr, 1); rows = {str(ws.cell(i, wc['Name']).value or '').strip(): i for i in range(2, ws.max_row + 1)}
    lg = wb['Additions Log'] if 'Additions Log' in wb.sheetnames else wb.create_sheet('Additions Log')
    for kept, loser, _, why in PAIRS: lg.append([today, 'merged', f'{loser} → {kept}', why])
    if 'Licence Watch' in wb.sheetnames:
        lw = wb['Licence Watch']
        for i in range(2, lw.max_row + 1):
            v = str(lw.cell(i, 1).value or '').strip()
            for kept, loser, _, _ in PAIRS:
                if v == loser: lw.cell(i, 1).value = kept
    wb.save(a.out or a.workbook); print('workbook saved:', a.out or a.workbook, '| rows', ws.max_row - 1)
