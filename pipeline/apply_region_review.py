#!/usr/bin/env python3
"""
apply_region_review.py — applies Youssef's decisions from the region-review workbook (Oct 2026) to the three places a
stakeholder lives: the stakeholder workbook, the hub's Stakeholder Map (hub/EH_Stakeholder_Map_CURRENT.html) and the
pipeline's durable tables (aliases.csv, region_overrides.csv).

What it does
  1. Regions — every row with Decision = Accept gets its Proposed Region written to the workbook (Region / City) and to the
     map record (rg / city). The same name → region pairs are written to pipeline/region_overrides.csv, which
     sync_bids_to_apps.py reads on every publish, so a region survives the sync even if the record is re-created.
  2. Merges — a row whose Twin Candidate names another record (and whose decision is Accept) is a duplicate: the duplicate
     record is removed from the workbook and the map, its name is recorded on the kept record as
     "Alias (merged duplicate): …" (what the map's search and the sync's matcher use), and an approved row
     tracker_name → map_name is appended to aliases.csv so the rosters resolve to the kept record from now on.
     Chains are followed (A → B → C ends at C) and a kept record is never itself removed.
  3. Rejected pairs — a Reject decision on a row with a twin candidate is written to aliases.csv as status=rejected, so no
     matcher pairs those two names again.
  4. Logs — a "Merge Log" and a "Region Log" sheet in the workbook, and a change table (MD + JSON) next to this script.

Usage:
  python pipeline/apply_region_review.py --review EH_Region_Review_Batch1.xlsx --workbook "EH_Stakeholder_Competitive_Map 4.xlsx" \
      --map hub/EH_Stakeholder_Map_CURRENT.html --out-workbook "EH_Stakeholder_Competitive_Map 5.xlsx" [--hold 4061,4078,...]
--hold lists Row IDs to leave untouched (decisions queried back to the reviewer). --dry-run prints and writes nothing.
"""
import argparse, csv, json, os, re, sys, datetime as dt
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from sync_bids_to_apps import norm

REGIONS = ['Riyadh', 'Eastern Province', 'Jeddah', 'Southern Region', 'Makkah City', 'Madinah Region', 'Qassim Region',
           'GCC / International', 'Other Makkah Region', 'Al Jawf Region', 'Tabuk Region', 'Hail Region', 'Rabigh', 'Northern Borders Region']
REGION_FIX = {'eastern region': 'Eastern Province', 'eastern': 'Eastern Province', 'riyadh region': 'Riyadh', 'jeddah city': 'Jeddah',
              'asir': 'Southern Region', 'makkah': 'Makkah City', 'madinah': 'Madinah Region', 'qassim': 'Qassim Region', 'tabuk': 'Tabuk Region', 'hail': 'Hail Region'}

def fix_region(v):
    v = str(v or '').strip()
    if v in REGIONS: return v
    return REGION_FIX.get(v.lower(), None)

def strip_twin(t):
    """'Name [Region] — note' → 'Name'; also 'Name [row 585]'."""
    t = str(t or '').strip()
    if not t or t in ('—', '-', 'nan'): return ''
    t = re.sub(r'\s*\[[^\]]*\]', '', t)
    t = t.split(' — ')[0].split(';')[0]
    return t.strip()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--review', required=True); ap.add_argument('--workbook', required=True); ap.add_argument('--map', required=True)
    ap.add_argument('--out-workbook', required=True); ap.add_argument('--out-map', default=None)
    ap.add_argument('--aliases', default=os.path.join(HERE, 'aliases.csv')); ap.add_argument('--overrides', default=os.path.join(HERE, 'region_overrides.csv'))
    ap.add_argument('--hold', default=''); ap.add_argument('--dry-run', action='store_true')
    ap.add_argument('--extra-merge', action='append', default=[], help='"loser => kept" pairs decided outside the review sheet')
    a = ap.parse_args()
    hold = {int(x) for x in a.hold.split(',') if x.strip()}
    stamp = dt.datetime.now(dt.timezone(dt.timedelta(hours=3)))
    stamp_s, stamp_m = stamp.strftime('%Y-%m-%d'), stamp.strftime('%b %Y')

    # ---- review decisions
    rv = openpyxl.load_workbook(a.review, read_only=True)['Region Review']
    rows = list(rv.iter_rows(values_only=True)); hdr = [str(h or '') for h in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}
    decisions = []
    for r in rows[1:]:
        if r[col['Row ID']] is None: continue
        d = {h: r[i] for h, i in col.items()}
        d['dec'] = str(d.get('Decision') or '').strip().lower()
        d['twin'] = strip_twin(d.get('Twin Candidate'))
        d['Row ID'] = int(d['Row ID'])
        decisions.append(d)

    # ---- map
    mh = open(a.map, encoding='utf-8').read()
    m = re.search(r'(<script id="DATA" type="application/json">)(.*?)(</script>)', mh, re.S)
    D = json.loads(m.group(2)); nodes = D['nodes']
    by_norm = {}
    for n in nodes: by_norm.setdefault(norm(n['n']), n)
    by_raw = {n['n'].strip(): n for n in nodes}
    from difflib import SequenceMatcher
    def node_of(name):
        """exact name → normalised name → the one record whose normalised name contains it → closest at ≥0.9"""
        name = str(name).strip()
        if name in by_raw: return by_raw[name]
        k = norm(name)
        if k in by_norm: return by_norm[k]
        cont = [n for n in nodes if k and k in norm(n['n'])]
        if len({id(x) for x in cont}) == 1: return cont[0]
        best = max(nodes, key=lambda x: SequenceMatcher(None, k, norm(x['n'])).ratio())
        return best if SequenceMatcher(None, k, norm(best['n'])).ratio() >= 0.9 else None

    # ---- workbook
    wb = openpyxl.load_workbook(a.workbook)
    ws = wb['Stakeholders']
    whdr = [str(c.value or '') for c in ws[1]]
    wc = {h: i + 1 for i, h in enumerate(whdr)}
    wb_rows, wb_raw = {}, {}   # norm(name) → row number; raw name → row number
    for i in range(2, ws.max_row + 1):
        nm = ws.cell(i, wc['Name']).value
        if nm: wb_rows.setdefault(norm(nm), i); wb_raw.setdefault(str(nm).strip(), i)
    def wrow(name):
        name = str(name).strip()
        return wb_raw.get(name) or wb_rows.get(norm(name))

    # ---- build the plan
    region_set, merges, rejects, problems = [], [], [], []
    for d in decisions:
        rid, name, dec, twin = d['Row ID'], str(d['Name']).strip(), d['dec'], d['twin']
        if rid in hold: continue
        if dec == 'reject':
            for t in re.split(r';\s*', str(d.get('Twin Candidate') or '')):
                t = strip_twin(t)
                if t and norm(t) != norm(name): rejects.append((name, t))
            continue
        if dec != 'accept': continue
        reg = fix_region(d.get('Proposed Region'))
        if d.get('Proposed Region') and str(d['Proposed Region']).strip() not in ('', 'None', 'Searched — not found') and reg is None:
            problems.append(f"row {rid} {name}: region '{d['Proposed Region']}' is not on the list"); continue
        if reg: region_set.append({'row': rid, 'name': name, 'region': reg, 'city': str(d.get('Raw City Found') or '').strip(),
                                   'confidence': str(d.get('Confidence') or '').strip() or 'Reviewer', 'source': str(d.get('Source(s)') or '').strip()[:300],
                                   'method': str(d.get('Method') or '').strip()})
        if twin and norm(twin) != norm(name): merges.append((name, twin, f'review row {rid}'))
    for em in a.extra_merge:
        l, k = [x.strip() for x in em.split('=>')]; merges.append((l, k, 'decided in chat'))

    # ---- resolve merge chains, protect kept records
    # resolve every kept name to a real map record first, then follow chains on exact record names
    resolved = []
    for l, k, why in merges:
        kn = node_of(k)
        if kn is None: problems.append(f"merge {l} → {k}: kept record not on the map"); continue
        resolved.append((l.strip(), kn['n'].strip(), why))
    target = {l: k for l, k, _ in resolved}
    def final(k, seen=()):
        k = str(k).strip()
        if k in target and k not in seen: return final(target[k], seen + (k,))
        return k
    plan = []
    for l, k, why in resolved:
        kept = final(k); kn = by_raw[kept]; ln = by_raw.get(l)
        if kept == l: continue
        plan.append({'loser': l, 'kept': kept, 'why': why, 'loser_on_map': ln is not None, 'loser_in_workbook': wrow(l) is not None,
                     'kept_region': kn.get('rg'), 'loser_region': (ln or {}).get('rg'), 'loser_bc': (ln or {}).get('bc')})
    losers = {p['loser'] for p in plan}
    for p in plan:
        if p['kept'] in losers: problems.append(f"chain error: {p['kept']} is both kept and removed")
    if problems:
        print('PROBLEMS:'); [print(' -', p) for p in problems]
        if not a.dry_run: sys.exit('fix the problems above first')

    # ---- apply regions (to the record that will survive)
    region_log = []
    for r in region_set:
        tgt_name = final(r['name'])
        n = by_raw.get(tgt_name) or node_of(tgt_name)
        if n is None: problems.append(f"region: {tgt_name} not on the map"); continue
        old = n.get('rg')
        if tgt_name != r['name'].strip() and old not in ('—', '', None) and old != r['region']:
            # the kept twin already has a region that disagrees with the reviewer's choice for the duplicate
            region_log.append({**r, 'name': n['n'], 'from': old, 'to': r['region'], 'via': r['name'], 'note': 'kept record re-regioned to the reviewer\'s decision'}); n['rg'] = r['region']
        elif old in ('—', '', None) or old != r['region']:
            region_log.append({**r, 'name': n['n'], 'from': old, 'to': r['region'], 'via': r['name'] if tgt_name != r['name'].strip() else '', 'note': ''}); n['rg'] = r['region']
        if r['city'] and not n.get('city'): n['city'] = r['city']
        if a.dry_run: continue
        wr = wrow(n['n'])
        if wr:
            ws.cell(wr, wc['Region']).value = r['region']
            if r['city'] and not ws.cell(wr, wc['City']).value: ws.cell(wr, wc['City']).value = r['city']

    # ---- apply merges on the map
    for p in plan:
        kn = by_raw[p['kept']]; ln = by_raw.get(p['loser'])
        notes = kn.get('notes') or ''
        if f"Alias (merged duplicate): {p['loser']}" not in notes: kn['notes'] = (notes + ' ' if notes else '') + f"Alias (merged duplicate): {p['loser']}"
        if ln is not None:
            if ln.get('city') and not kn.get('city'): kn['city'] = ln['city']
            if (ln.get('bc') or 0) > (kn.get('bc') or 0):   # until the next sync recomputes it from the rosters
                kn['bid'] = True; kn['bc'] = ln['bc']; kn['bf'] = ln.get('bf'); kn['h2h'] = ln.get('h2h') or kn.get('h2h')
    nodes[:] = [n for n in nodes if n['n'].strip() not in losers]
    D['nodes'] = nodes

    # ---- apply merges in the workbook: Notes on the kept row, delete the duplicate row
    merge_log = []
    for p in plan:
        kr = wrow(p['kept']); lr = wb_raw.get(p['loser']) or (wb_rows.get(norm(p['loser'])) if norm(p['loser']) != norm(p['kept']) else None)
        merge_log.append({**p, 'kept_row': kr, 'loser_row': lr})
        if a.dry_run or kr is None: continue
        c = ws.cell(kr, wc['Notes']); old = c.value or ''
        if f"Alias (merged duplicate): {p['loser']}" not in old: c.value = (old + ' ' if old else '') + f"Alias (merged duplicate): {p['loser']}"
        if lr is not None:
            lb = ws.cell(lr, wc['Bid Encounters']).value or 0; kb = ws.cell(kr, wc['Bid Encounters']).value or 0
            if isinstance(lb, (int, float)) and isinstance(kb, (int, float)) and lb > kb:
                ws.cell(kr, wc['Bid Encounters']).value = lb; ws.cell(kr, wc['Bids vs EH']).value = 'Yes'
    del_rows = sorted({p['loser_row'] for p in merge_log if p['loser_row'] and p['loser_row'] != p['kept_row']}, reverse=True)
    if not a.dry_run:
        for r in del_rows: ws.delete_rows(r, 1)
        last = ws.max_row
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{last}"
        # Summary sheet COUNTIF ranges are fixed ($A$2:$A$4197): point them at the new last row
        if 'Summary' in wb.sheetnames:
            for row in wb['Summary'].iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith('='): c.value = re.sub(r'\$(\d{4})\b', lambda mm: f'${last}' if int(mm.group(1)) >= 4000 else mm.group(0), c.value)
        # Bid Rivals: regions follow; merged names are dropped (the sheet is regenerated by the pipeline anyway)
        if 'Bid Rivals' in wb.sheetnames:
            br = wb['Bid Rivals']; bh = [str(c.value or '') for c in br[1]]; bc = {h: i + 1 for i, h in enumerate(bh)}
            reg_by = {norm(r['name']): r['region'] for r in region_set}
            for i in range(br.max_row, 1, -1):
                nm = br.cell(i, bc['Name']).value
                if not nm: continue
                if str(nm).strip() in losers: br.delete_rows(i, 1); continue
                if norm(nm) in reg_by and 'Region' in bc: br.cell(i, bc['Region']).value = reg_by[norm(nm)]
        # logs
        for title, data, cols in (('Merge Log', merge_log, ['loser', 'kept', 'why', 'loser_row', 'kept_row', 'loser_region', 'kept_region', 'loser_bc']),
                                  ('Region Log', region_log, ['name', 'from', 'to', 'via', 'confidence', 'method', 'source', 'note'])):
            lg = wb[title] if title in wb.sheetnames else wb.create_sheet(title)
            if lg.max_row == 1 and lg.cell(1, 1).value is None:
                lg.append(['Date'] + cols)
                for c in lg[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E5F')
            for d in data: lg.append([stamp_s] + [d.get(k) for k in cols])
            for i, w in enumerate([11] + [40] * len(cols), 1): lg.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        wb.save(a.out_workbook)
        payload = json.dumps(D, ensure_ascii=False).replace('</', '<\\/')
        open(a.out_map or a.map, 'w', encoding='utf-8').write(mh[:m.start(2)] + payload + mh[m.end(2):])
        # aliases.csv
        fields = ['tracker_name', 'map_name', 'method', 'status', 'added']
        existing = list(csv.DictReader(open(a.aliases, encoding='utf-8-sig', newline=''))) if os.path.exists(a.aliases) else []
        known = {(norm(r['tracker_name']), norm(r['map_name'])) for r in existing}
        new = []
        for p in plan:
            if (norm(p['loser']), norm(p['kept'])) not in known: new.append({'tracker_name': p['loser'], 'map_name': p['kept'], 'method': 'reviewer merge (region review)', 'status': 'approved', 'added': stamp_m})
        for l, t in rejects:
            if (norm(l), norm(t)) not in known: new.append({'tracker_name': l, 'map_name': t, 'method': 'reviewer: different companies', 'status': 'rejected', 'added': stamp_m})
        with open(a.aliases, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
            for r in existing + new: w.writerow({k: r.get(k, '') for k in fields})
        # region_overrides.csv (keyed by name; the sync applies it to any record whose region is '—')
        ofields = ['map_name', 'region', 'city', 'confidence', 'method', 'source', 'decided']
        oex = list(csv.DictReader(open(a.overrides, encoding='utf-8-sig', newline=''))) if os.path.exists(a.overrides) else []
        okeys = {norm(r['map_name']) for r in oex}
        for r in region_set:
            for nm in {r['name'], final(r['name'])}:
                if norm(nm) in okeys: continue
                okeys.add(norm(nm)); oex.append({'map_name': nm, 'region': r['region'], 'city': r['city'], 'confidence': r['confidence'], 'method': r['method'], 'source': r['source'], 'decided': stamp_s})
        with open(a.overrides, 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.DictWriter(f, fieldnames=ofields); w.writeheader(); [w.writerow({k: r.get(k, '') for k in ofields}) for r in oex]
        # change table
        ct = {'date': stamp_s, 'regions_set': region_log, 'merges': merge_log, 'rejected_pairs': rejects, 'held_rows': sorted(hold), 'workbook_rows_deleted': len(del_rows), 'map_records_removed': len(losers)}
        json.dump(ct, open(os.path.join(HERE, 'change_table_region_review.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        with open(os.path.join(HERE, 'change_table_region_review.md'), 'w', encoding='utf-8') as f:
            f.write(f"# Region review — applied {stamp_s}\n\n## Regions set ({len(region_log)})\n| Record | From | To | Via duplicate | Confidence |\n|---|---|---|---|---|\n")
            for r in region_log: f.write(f"| {r['name']} | {r['from']} | {r['to']} | {r['via']} | {r['confidence']} |\n")
            f.write(f"\n## Merges ({len(merge_log)})\n| Removed (duplicate) | Kept record | Basis |\n|---|---|---|\n")
            for p in merge_log: f.write(f"| {p['loser']} | {p['kept']} | {p['why']} |\n")
            f.write(f"\n## Rejected pairs ({len(rejects)})\n")
            for l, t in rejects: f.write(f"- {l} ≠ {t}\n")
            f.write(f"\n## Held for confirmation: {sorted(hold)}\n")
    print(f"regions set: {len(region_log)} · merges: {len(plan)} · workbook rows deleted: {len(del_rows) if not a.dry_run else '(dry run)'} · rejected pairs: {len(rejects)} · held: {sorted(hold)}")
    for r in region_log: print(f"  region  {r['name']}: {r['from']} → {r['to']}" + (f"  (via {r['via']})" if r['via'] else ''))
    for p in plan: print(f"  merge   {p['loser']}  →  {p['kept']}   [{p['why']}]" + ('' if p['loser_on_map'] else '  (loser not on map)'))

if __name__ == '__main__':
    main()
