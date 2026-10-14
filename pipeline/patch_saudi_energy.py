"""patch_saudi_energy.py — Saudi Electricity Company (SEC) rebranded as Saudi Energy (الشركة السعودية للطاقة) in Feb 2026;
the hub held it as two records. One record now, everywhere (14 Oct 2026).

Stakeholder Map (hub/EH_Stakeholder_Map_CURRENT.html):
  • "Saudi Electricity Company (SEC)" (EH Clients, Tier 1, score 100) is kept and renamed "Saudi Energy (formerly SEC)"
  • "Saudi Energy" (Clients / Potential Clients, score 0 — the duplicate) is removed; its HQ pin is dropped
  • every old spelling is recorded on the kept record as "Alias (merged duplicate): …" so search and the bid sync still find it
  • the HQ-pin table and the Bid View block (BANALYTICS) use the new name (BANALYTICS is also rebuilt hourly from the bid app)
  • "Saudi Electricity Project Development Co." is a separate subsidiary and is NOT folded
Client Bubble Map (hub/EH_Client_Bubble_Map_CURRENT.html): the client «الشركة السعودية للكهرباء (SE)» is shown as
  «الشركة السعودية للطاقة (SE)» — same figures, one bubble.
Stakeholder workbook (optional --workbook IN --out OUT): the same rename and removal, logged in the Merge Log sheet.
Run once: python3 pipeline/patch_saudi_energy.py [--workbook IN.xlsx --out OUT.xlsx]
"""
import argparse, json, os, re, sys, datetime as dt
HERE = os.path.dirname(os.path.abspath(__file__))
OLD, DUP, NEW = 'Saudi Electricity Company (SEC)', 'Saudi Energy', 'Saudi Energy (formerly SEC)'
ALIASES = ['Saudi Electricity Company (SEC)', 'Saudi Electricity Company', 'Saudi Energy', 'الشركة السعودية للطاقة', 'الشركة السعودية للكهرباء', 'SEC']
NOTE = ('Rebranded from Saudi Electricity Company (SEC) to Saudi Energy (الشركة السعودية للطاقة), Feb 2026. '
        'Alias (merged duplicate): ' + '; '.join(ALIASES))

ap = argparse.ArgumentParser(); ap.add_argument('--map', default=os.path.join(HERE, '..', 'hub', 'EH_Stakeholder_Map_CURRENT.html'))
ap.add_argument('--workbook'); ap.add_argument('--out'); a = ap.parse_args()
stamp = dt.datetime.now(dt.timezone(dt.timedelta(hours=3))).strftime('%Y-%m-%d')

h = open(a.map, encoding='utf-8').read()
m = re.search(r'(<script id="DATA" type="application/json">)(.*?)(</script>)', h, re.S)
D = json.loads(m.group(2)); nodes = D['nodes']
kept = next((n for n in nodes if n['n'] in (OLD, NEW)), None)
dup = next((n for n in nodes if n['n'] == DUP), None)
if kept is None: sys.exit('kept record not found')
changed = []
if kept['n'] == OLD:
    kept['n'] = NEW; changed.append(f'renamed {OLD} → {NEW}')
if 'Rebranded from Saudi Electricity Company' not in (kept.get('notes') or ''):
    kept['notes'] = ((kept.get('notes') or '') + ' ' + NOTE).strip(); changed.append('aliases recorded')
if dup is not None:
    for f in ('city', 'reg'):
        if dup.get(f) and not kept.get(f): kept[f] = dup[f]
    if (dup.get('bc') or 0) > (kept.get('bc') or 0): kept.update(bid=dup['bid'], bc=dup['bc'], bf=dup.get('bf'), h2h=dup.get('h2h'))
    nodes.remove(dup); changed.append(f'removed duplicate "{DUP}"')
payload = json.dumps(D, ensure_ascii=False).replace('</', '<\\/')
h = h[:m.start(2)] + payload + h[m.end(2):]
# outside DATA: HQ-pin table and Bid View block
n_pin = len(re.findall(r'"Saudi Energy":\[[^\]]*\],', h))
h = re.sub(r'"Saudi Energy":\[[^\]]*\],', '', h)
n_old = h.count(f'"{OLD}"'); h = h.replace(f'"{OLD}"', f'"{NEW}"')
if n_pin or n_old: changed.append(f'{n_old} references renamed outside the data block, {n_pin} duplicate pin removed')
open(a.map, 'w', encoding='utf-8').write(h)
print('map:', '; '.join(changed) or 'already applied')

B = os.path.join(HERE, '..', 'hub', 'EH_Client_Bubble_Map_CURRENT.html')
if os.path.exists(B):
    b = open(B, encoding='utf-8').read(); nb = b.count('الشركة السعودية للكهرباء (SE)')
    if nb: open(B, 'w', encoding='utf-8').write(b.replace('الشركة السعودية للكهرباء (SE)', 'الشركة السعودية للطاقة (SE)'))
    print('bubble map:', f'{nb} references renamed' if nb else 'already applied')

if a.workbook:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.load_workbook(a.workbook); ws = wb['Stakeholders']
    hdr = {str(c.value or ''): i + 1 for i, c in enumerate(ws[1])}
    kr = dr = None
    for i in range(2, ws.max_row + 1):
        v = str(ws.cell(i, hdr['Name']).value or '').strip()
        if v in (OLD, NEW): kr = i
        elif v == DUP: dr = i
    if kr is None: sys.exit('workbook: kept row not found')
    ws.cell(kr, hdr['Name']).value = NEW
    c = ws.cell(kr, hdr['Notes'])
    if 'Rebranded from Saudi Electricity Company' not in str(c.value or ''): c.value = (str(c.value or '') + ' ' + NOTE).strip()
    if dr:
        ws.delete_rows(dr, 1)
        lg = wb['Merge Log'] if 'Merge Log' in wb.sheetnames else wb.create_sheet('Merge Log')
        if lg.max_row == 1 and lg.cell(1, 1).value is None:
            lg.append(['Date', 'loser', 'kept', 'why'])
            for x in lg[1]: x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E5F')
        lg.append([stamp, DUP, NEW, 'SEC rebranded as Saudi Energy (Feb 2026) — same company; kept record renamed from ' + OLD, dr, kr])
        last = ws.max_row
        if 'Summary' in wb.sheetnames:
            for row in wb['Summary'].iter_rows():
                for x in row:
                    if isinstance(x.value, str) and x.value.startswith('='): x.value = re.sub(r'\$(\d{4})\b', lambda mm: f'${last}' if int(mm.group(1)) >= 4000 else mm.group(0), x.value)
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{last}"
    for sh in wb.sheetnames:   # any other sheet that names the old record
        if sh in ('Stakeholders', 'Merge Log', 'Region Log'): continue
        for row in wb[sh].iter_rows():
            for x in row:
                if x.value == OLD: x.value = NEW
    wb.save(a.out or a.workbook); print('workbook: renamed row', kr, '; removed duplicate row', dr)
